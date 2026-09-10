#!/usr/bin/env python3
"""Презентация: Острая травма зубов (БГМУ, 5 курс) #5"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Цветовая схема — медицинская, спокойная (не фиолетовая)
NAVY = RGBColor(0x1A, 0x3A, 0x4A)
TEAL = RGBColor(0x2A, 0x6F, 0x7A)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)
LIGHT_BG = RGBColor(0xF5, 0xF7, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT = RGBColor(0x1E, 0x2A, 0x32)
MUTED = RGBColor(0x5A, 0x6B, 0x75)
SOFT_TEAL = RGBColor(0xE8, 0xF1, 0xF3)
SOFT_ACCENT = RGBColor(0xFB, 0xEE, 0xE6)


def set_run_font(run, size=18, bold=False, color=DARK_TEXT, font_name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name


def add_bg_rect(slide, prs, color):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    return shape


def add_top_bar(slide, prs, color=TEAL, height=0.12):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(height)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    return bar


def add_footer(slide, prs, text="БГМУ · 5 курс · Острая травма зубов · #5", page=None, total=None):
    box = slide.shapes.add_textbox(
        Inches(0.5), Inches(6.95), Inches(8.5), Inches(0.35)
    )
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_font(run, size=11, color=MUTED)
    if page is not None and total is not None:
        num = slide.shapes.add_textbox(
            Inches(11.5), Inches(6.95), Inches(1.3), Inches(0.35)
        )
        ntf = num.text_frame
        np = ntf.paragraphs[0]
        np.alignment = PP_ALIGN.RIGHT
        nrun = np.add_run()
        nrun.text = f"{page} / {total}"
        set_run_font(nrun, size=11, color=MUTED)


def add_title(slide, text, left=0.5, top=0.35, width=12.3, height=0.7, size=28, color=NAVY):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=True, color=color)
    return box


def add_bullets(slide, items, left=0.5, top=1.2, width=12.3, height=5.2, size=18, spacing=1.15):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        # item can be str or (str, level) or (str, bold)
        if isinstance(item, tuple):
            text, extra = item
            if isinstance(extra, int):
                p.level = extra
                run = p.add_run()
                run.text = ("• " if p.level == 0 else "– ") + text
                set_run_font(run, size=size - (2 if p.level else 0), color=DARK_TEXT)
            elif extra == "bold":
                run = p.add_run()
                run.text = "• " + text
                set_run_font(run, size=size, bold=True, color=NAVY)
            else:
                run = p.add_run()
                run.text = "• " + text
                set_run_font(run, size=size, color=DARK_TEXT)
        else:
            run = p.add_run()
            run.text = "• " + item
            set_run_font(run, size=size, color=DARK_TEXT)
    return box


def add_card(slide, left, top, width, height, title, body_lines, fill=SOFT_TEAL, title_color=TEAL):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    card.line.fill.background()
    # title
    tbox = slide.shapes.add_textbox(
        Inches(left + 0.2), Inches(top + 0.15), Inches(width - 0.4), Inches(0.4)
    )
    tp = tbox.text_frame.paragraphs[0]
    tr = tp.add_run()
    tr.text = title
    set_run_font(tr, size=16, bold=True, color=title_color)
    # body
    bbox = slide.shapes.add_textbox(
        Inches(left + 0.2), Inches(top + 0.55), Inches(width - 0.4), Inches(height - 0.7)
    )
    btf = bbox.text_frame
    btf.word_wrap = True
    for i, line in enumerate(body_lines):
        p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = line
        set_run_font(r, size=13, color=DARK_TEXT)


def make_prs():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slides_meta = []  # for page numbers later

    # ========== 1. ТИТУЛЬНЫЙ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, NAVY)
    # accent strip
    strip = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(2.4), prs.slide_width, Inches(2.4)
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = TEAL
    strip.line.fill.background()

    box = s.shapes.add_textbox(Inches(0.8), Inches(2.55), Inches(11.5), Inches(1.2))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "Острая травма зубов"
    set_run_font(r, size=44, bold=True, color=WHITE)

    sub = s.shapes.add_textbox(Inches(0.8), Inches(3.7), Inches(11.5), Inches(0.6))
    sp = sub.text_frame.paragraphs[0]
    sr = sp.add_run()
    sr.text = "Презентация · Белорусский государственный медицинский университет"
    set_run_font(sr, size=18, color=RGBColor(0xD0, 0xE4, 0xE8))

    meta = s.shapes.add_textbox(Inches(0.8), Inches(5.6), Inches(11.5), Inches(0.8))
    mp = meta.text_frame.paragraphs[0]
    mr = mp.add_run()
    mr.text = "5 курс  ·  Стоматологический факультет  ·  Лекция №5"
    set_run_font(mr, size=16, color=RGBColor(0xA8, 0xC0, 0xC8))
    slides_meta.append(s)

    # ========== 2. ПЛАН ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "План лекции")
    add_bullets(
        s,
        [
            "Определение и эпидемиология",
            "Классификация травм зубов (ВОЗ / Andreasen)",
            "Обследование пациента с травмой",
            "Травмы твёрдых тканей коронки и корня",
            "Травмы периодонта (люксации, вывих, полный вывих)",
            "Неотложная помощь и лечение",
            "Осложнения и диспансеризация",
            "Особенности травмы временных зубов",
        ],
        top=1.15,
        size=20,
    )
    slides_meta.append(s)

    # ========== 3. ОПРЕДЕЛЕНИЕ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Определение")
    add_bullets(
        s,
        [
            ("Острая травма зуба — повреждение зуба и/или окружающих тканей "
             "вследствие одномоментного воздействия внешней силы.", "bold"),
            "Включает повреждения эмали, дентина, пульпы, цемента, периодонта, альвеолярной кости и мягких тканей.",
            "Чаще всего страдают передние зубы верхней челюсти (центральные резцы).",
            "Пик травматизма: 2–4 года (временные) и 8–12 лет (постоянные).",
            "Основные причины: падения, спорт, ДТП, бытовой и криминальный травматизм.",
        ],
        top=1.15,
        size=18,
    )
    slides_meta.append(s)

    # ========== 4. ЭПИДЕМИОЛОГИЯ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Эпидемиология и факторы риска")
    add_card(
        s, 0.5, 1.2, 3.9, 4.8,
        "Частота",
        [
            "До 25–30% детей имеют",
            "опыт травмы постоянных",
            "зубов к 14 годам",
            "",
            "Мальчики травмируются",
            "примерно в 1,5–2 раза чаще",
            "",
            "Верхние центральные",
            "резцы — до 70–80%",
            "всех случаев",
        ],
        SOFT_TEAL, TEAL,
    )
    add_card(
        s, 4.7, 1.2, 3.9, 4.8,
        "Факторы риска",
        [
            "Протрузия резцов,",
            "открытый прикус",
            "",
            "Недостаточная длина",
            "верхней губы",
            "",
            "Кариес, ослабленная",
            "коронка",
            "",
            "Эпилепсия, ADHD,",
            "нарушения координации",
        ],
        SOFT_ACCENT, ACCENT,
    )
    add_card(
        s, 8.9, 1.2, 3.9, 4.8,
        "Ситуации",
        [
            "Спорт (хоккей, футбол,",
            "велосипед, скейт)",
            "",
            "Игровая площадка,",
            "школа",
            "",
            "ДТП и бытовые",
            "несчастные случаи",
            "",
            "Драка / насилие",
            "(исключить жестокость",
            "к ребёнку!)",
        ],
        SOFT_TEAL, TEAL,
    )
    slides_meta.append(s)

    # ========== 5. КЛАССИФИКАЦИЯ ОБЗОР ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Классификация (ВОЗ / Andreasen)")
    add_bullets(
        s,
        [
            ("I. Травмы твёрдых тканей зуба и пульпы", "bold"),
            ("Трещина эмали; перелом коронки (неосложнённый / осложнённый); перелом коронки и корня; перелом корня", 1),
            ("II. Травмы периодонта", "bold"),
            ("Сотрясение; подвывих; экструзивный, латеральный, интрузивный вывих; полный вывих (авульсия)", 1),
            ("III. Травмы поддерживающей кости", "bold"),
            ("Перелом стенки альвеолы, альвеолярного отростка, челюсти", 1),
            ("IV. Травмы мягких тканей", "bold"),
            ("Ушиб, ссадина, разрыв, укус губ, десны, языка, слизистой", 1),
        ],
        top=1.15,
        size=17,
    )
    slides_meta.append(s)

    # ========== 6. ОБСЛЕДОВАНИЕ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Обследование пациента")
    add_bullets(
        s,
        [
            ("Анамнез:", "bold"),
            ("Когда, где, как произошла травма; тетанус; потеря сознания / тошнота (исключить ЧМТ!)", 1),
            ("Клиника:", "bold"),
            ("Осмотр лица, мягких тканей, прикуса; подвижность; перкуссия; зондирование; цвет коронки", 1),
            ("Тесты пульпы:", "bold"),
            ("ЭОД / холодовая проба — осторожно: сразу после травмы возможна ложная отрицательная реакция", 1),
            ("Рентген:", "bold"),
            ("Прицельный снимок, окклюзионный, при необходимости КЛКТ; поиск фрагментов в мягких тканях", 1),
            ("Фотофиксация и согласие на лечение — обязательны при травме", 1),
        ],
        top=1.1,
        size=17,
    )
    slides_meta.append(s)

    # ========== 7. ТРЕЩИНА И ПЕРЕЛОМ КОРОНКИ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Трещина эмали и перелом коронки")
    add_card(
        s, 0.5, 1.2, 6.0, 5.0,
        "Неосложнённый перелом",
        [
            "• Только эмаль или эмаль + дентин",
            "  без вскрытия пульпы",
            "",
            "• Лечение: сглаживание краёв,",
            "  покрытие дентина (ГИС / адгезив),",
            "  реставрация композитом",
            "  или реаттачмент фрагмента",
            "",
            "• Контроль витальности пульпы",
            "  через 1, 3, 6, 12 мес.",
        ],
        SOFT_TEAL, TEAL,
    )
    add_card(
        s, 6.8, 1.2, 6.0, 5.0,
        "Осложнённый перелом",
        [
            "• Вскрытие пульпы",
            "",
            "• Несформированный корень:",
            "  прямое покрытие / частичная",
            "  пульпотомия (Cvek) — MTA / BC",
            "",
            "• Сформированный корень:",
            "  эндодонтическое лечение",
            "  или витальные методы",
            "  по показаниям",
            "",
            "• Затем эстетическая реставрация",
        ],
        SOFT_ACCENT, ACCENT,
    )
    slides_meta.append(s)

    # ========== 8. ПЕРЕЛОМ КОРНЯ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Перелом корня")
    add_bullets(
        s,
        [
            "Классификация по уровню: апикальная, средняя, корональная треть.",
            "Клиника: подвижность коронковой части, болезненная перкуссия, иногда смещение.",
            "Диагностика: несколько прицельных снимков под разными углами / КЛКТ.",
            ("Лечение:", "bold"),
            ("Репозиция коронкового фрагмента и шинирование на 4 недели (при корональной трети — дольше).", 1),
            ("Контроль пульпы; при некрозе — эндодонтия только коронкового фрагмента.", 1),
            ("Апикальный фрагмент часто остаётся витальным и не требует вмешательства.", 1),
            "Прогноз лучше при апикальных переломах и несформированной верхушке.",
        ],
        top=1.15,
        size=17,
    )
    slides_meta.append(s)

    # ========== 9. ЛЮКСАЦИИ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Травмы периодонта (люксации)")
    items = [
        ("Сотрясение", "Боль при перкуссии, без смещения и подвижности. Наблюдение."),
        ("Подвывих", "Повышенная подвижность без смещения. Шинирование 2 нед. при необходимости."),
        ("Экструзия", "Частичное выдвижение из лунки по оси. Репозиция + шина 2 нед."),
        ("Латеральный вывих", "Смещение с переломом/сдавлением альвеолы. Репозиция + шина 4 нед."),
        ("Интрузия", "Вколоченный зуб. Тактика зависит от степени и стадии корня."),
    ]
    y = 1.15
    for title, desc in items:
        row = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(y), Inches(12.3), Inches(0.95)
        )
        row.fill.solid()
        row.fill.fore_color.rgb = WHITE
        row.line.color.rgb = RGBColor(0xD0, 0xDC, 0xE0)
        tbox = s.shapes.add_textbox(Inches(0.7), Inches(y + 0.12), Inches(3.2), Inches(0.7))
        tp = tbox.text_frame.paragraphs[0]
        tr = tp.add_run()
        tr.text = title
        set_run_font(tr, size=15, bold=True, color=TEAL)
        dbox = s.shapes.add_textbox(Inches(4.0), Inches(y + 0.12), Inches(8.5), Inches(0.7))
        dtf = dbox.text_frame
        dtf.word_wrap = True
        dp = dtf.paragraphs[0]
        dr = dp.add_run()
        dr.text = desc
        set_run_font(dr, size=14, color=DARK_TEXT)
        y += 1.05
    slides_meta.append(s)

    # ========== 10. ИНТРУЗИЯ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Интрузивный вывих — тактика")
    add_bullets(
        s,
        [
            ("Несформированный корень (открытая верхушка):", "bold"),
            ("Допускается выжидание спонтанной реэрупции (до 3–4 недель наблюдения).", 1),
            ("При отсутствии движения — ортодонтическая или хирургическая репозиция.", 1),
            ("Сформированный корень:", "bold"),
            ("Хирургическая или ортодонтическая репозиция + шинирование 4 недели.", 1),
            ("Высокий риск некроза пульпы → ранняя эндодонтия, профилактика резорбции (гидроксид кальция / MTA).", 1),
            "Риск: заменительная резорбция корня, анкилоз, потеря зуба.",
            "Контроль: рентген и клинический осмотр по графику IADT.",
        ],
        top=1.15,
        size=17,
    )
    slides_meta.append(s)

    # ========== 11. АВУЛЬСИЯ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Полный вывих (авульсия) — «золотой час»")
    add_card(
        s, 0.5, 1.15, 6.0, 5.1,
        "Первая помощь на месте",
        [
            "1. Найти зуб, держать за коронку!",
            "2. Не тереть корень, не scrub",
            "3. При загрязнении — промыть",
            "   физраствором / молоком",
            "4. По возможности — немедленная",
            "   реплантация в лунку",
            "5. Если нельзя — хранить в:",
            "   • молоке",
            "   • физрастворе / HBSS",
            "   • слюне (щечный карман)",
            "6. НЕ в воде! НЕ в сухую!",
            "7. Срочно к стоматологу",
        ],
        SOFT_ACCENT, ACCENT,
    )
    add_card(
        s, 6.8, 1.15, 6.0, 5.1,
        "В клинике",
        [
            "• Оценить время внелунки",
            "  и среду хранения",
            "• Промыть лунку, реплантировать",
            "• Гибкая шина 2 недели",
            "  (при переломе кости — 4 нед.)",
            "• Антибиотик, столбняк —",
            "  по показаниям",
            "• Открытая верхушка: возможна",
            "  реваскуляризация",
            "• Закрытая верхушка: эндодонтия",
            "  через 7–10 дней",
            "• Прогноз ↓ при >60 мин сухого",
            "  хранения",
        ],
        SOFT_TEAL, TEAL,
    )
    slides_meta.append(s)

    # ========== 12. ШИНИРОВАНИЕ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Шинирование")
    add_bullets(
        s,
        [
            "Предпочтительны гибкие шины (проволока + композит, стекловолокно) — сохраняют физиологическую подвижность.",
            "Жёсткие шины и длительная иммобилизация повышают риск анкилоза и резорбции.",
            ("Ориентиры по срокам (IADT):", "bold"),
            ("Подвывих / экструзия / авульсия — около 2 недель", 1),
            ("Латеральный вывих / интрузия / перелом корня — около 4 недель", 1),
            ("Перелом альвеолы — до 4 недель", 1),
            "Гигиена, мягкая диета, исключение нагрузки на травмированные зубы.",
            "Контроль окклюзии: устранить преждевременные контакты.",
        ],
        top=1.15,
        size=17,
    )
    slides_meta.append(s)

    # ========== 13. ВРЕМЕННЫЕ ЗУБЫ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Особенности травмы временных зубов")
    add_bullets(
        s,
        [
            "Главный принцип: не навредить зачатку постоянного зуба.",
            "Полный вывих временного зуба — НЕ реплантируют.",
            "Интрузия: часто выжидательная тактика; при направлении на зачаток — удаление.",
            "Переломы корня / выраженное смещение — чаще удаление коронкового фрагмента.",
            "Родители: предупредить о возможных нарушениях прорезывания постоянного зуба",
            "(гипоплазия эмали, дистопия, задержка прорезывания, редко — одонтома).",
            "Обязателен рентген-контроль и наблюдение до прорезывания постоянного зуба.",
        ],
        top=1.15,
        size=18,
    )
    slides_meta.append(s)

    # ========== 14. ОСЛОЖНЕНИЯ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Осложнения")
    add_card(
        s, 0.5, 1.2, 4.0, 5.0,
        "Пульпа",
        [
            "• Некроз пульпы",
            "• Облитерация",
            "  корневого канала",
            "• Внутренняя",
            "  резорбция",
            "• Периапикальный",
            "  периодонтит",
        ],
        SOFT_TEAL, TEAL,
    )
    add_card(
        s, 4.7, 1.2, 4.0, 5.0,
        "Периодонт / корень",
        [
            "• Воспалительная",
            "  резорбция корня",
            "• Заменительная",
            "  резорбция",
            "  (анкилоз)",
            "• Потеря зуба",
            "• Нарушение роста",
            "  альвеолы",
        ],
        SOFT_ACCENT, ACCENT,
    )
    add_card(
        s, 8.9, 1.2, 4.0, 5.0,
        "Прочее",
        [
            "• Дисколорит",
            "  коронки",
            "• Свищ, абсцесс",
            "• Эстетические",
            "  дефекты",
            "• Психологический",
            "  дискомфорт",
            "  пациента",
        ],
        SOFT_TEAL, TEAL,
    )
    slides_meta.append(s)

    # ========== 15. ДИСПАНСЕРИЗАЦИЯ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Диспансеризация (ориентир IADT)")
    add_bullets(
        s,
        [
            "Типичный график контроля: 2 недели → 4 недели → 3 месяца → 6 месяцев → 1 год → далее ежегодно.",
            "На каждом визите: жалобы, цвет, подвижность, перкуссия, зондирование десны, тест пульпы, рентген.",
            "Ранние признаки неблагополучия: серый/розовый цвет, свищ, нарастающая подвижность, просветление у верхушки,",
            "симптомы резорбции на снимке.",
            "Документирование динамики обязательно (карта травмы, фото, снимки).",
            "При несформированных корнях — особое внимание к продолжению апексогенеза.",
        ],
        top=1.15,
        size=17,
    )
    slides_meta.append(s)

    # ========== 16. ПРОФИЛАКТИКА ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Профилактика")
    add_bullets(
        s,
        [
            "Индивидуальные капы при контактных видах спорта.",
            "Раннее ортодонтическое лечение выраженной протрузии резцов.",
            "Безопасная среда: ремни безопасности, шлемы, ограждения площадок.",
            "Обучение населения и педагогов алгоритму первой помощи при авульсии.",
            "Санация полости рта — снижение риска «слабой» коронки при ударе.",
            "Информирование родителей дошкольников о высоком риске травм в 2–4 года.",
        ],
        top=1.15,
        size=18,
    )
    slides_meta.append(s)

    # ========== 17. КЛЮЧЕВЫЕ ВЫВОДЫ ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, LIGHT_BG)
    add_top_bar(s, prs)
    add_title(s, "Ключевые выводы")
    add_bullets(
        s,
        [
            "Травма зуба — неотложное состояние: время решает прогноз, особенно при авульсии.",
            "Диагноз строится на клинике + рентгене; тест пульпы сразу после травмы ненадёжен.",
            "Гибкое краткосрочное шинирование и контроль окклюзии — основа лечения люксаций.",
            "Временные зубы: приоритет — сохранность зачатка постоянного; авульсию не реплантируют.",
            "Длительное наблюдение необходимо для раннего выявления некроза и резорбции.",
            "Следуйте актуальным рекомендациям IADT и национальным протоколам.",
        ],
        top=1.15,
        size=18,
    )
    slides_meta.append(s)

    # ========== 18. СПАСИБО ==========
    s = prs.slides.add_slide(blank)
    add_bg_rect(s, prs, NAVY)
    strip = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(2.6), prs.slide_width, Inches(2.0)
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = TEAL
    strip.line.fill.background()

    box = s.shapes.add_textbox(Inches(0.8), Inches(2.85), Inches(11.5), Inches(0.8))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Спасибо за внимание"
    set_run_font(r, size=40, bold=True, color=WHITE)

    sub = s.shapes.add_textbox(Inches(0.8), Inches(3.7), Inches(11.5), Inches(0.5))
    sp = sub.text_frame.paragraphs[0]
    sp.alignment = PP_ALIGN.CENTER
    sr = sp.add_run()
    sr.text = "Вопросы?  ·  БГМУ, 5 курс  ·  Лекция №5"
    set_run_font(sr, size=18, color=RGBColor(0xD0, 0xE4, 0xE8))

    lit = s.shapes.add_textbox(Inches(0.8), Inches(5.5), Inches(11.5), Inches(1.2))
    lp = lit.text_frame
    lp.word_wrap = True
    p1 = lp.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = "Литература: рекомендации IADT (Dental Traumatology); клинические протоколы РБ;"
    set_run_font(r1, size=13, color=RGBColor(0xA8, 0xC0, 0xC8))
    p2 = lp.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "Andreasen J.O. et al. Traumatic Dental Injuries; учебники терапевтической стоматологии БГМУ."
    set_run_font(r2, size=13, color=RGBColor(0xA8, 0xC0, 0xC8))
    slides_meta.append(s)

    # footers (skip title and thanks)
    total = len(slides_meta)
    for i, slide in enumerate(slides_meta):
        if i == 0 or i == total - 1:
            continue
        add_footer(slide, prs, page=i + 1, total=total)

    out = "/workspace/Острая_травма_зубов_БГМУ_5курс.pptx"
    prs.save(out)
    return out, total


if __name__ == "__main__":
    path, n = make_prs()
    print(f"Saved: {path} ({n} slides)")
